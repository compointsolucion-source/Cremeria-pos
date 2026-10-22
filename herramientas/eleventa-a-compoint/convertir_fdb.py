#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convierte la base de datos de Eleventa (PDVDATA.FDB) en Excel listos para Compoint Punto.

Saca SOLO información del propio negocio:
  - Productos: código, nombre, costo, precio, mayoreo, existencia, mínimo, máximo,
    departamento y si se vende A Granel o Por Pieza.
  - Clientes: nombre, teléfono, límite de crédito y lo que deben.

NO toca el "catálogo base" que trae Eleventa (tablas PRODUCTOS_BASE y PROVEEDORES_BASE):
ese catálogo es de Eleventa y viene protegido.

Resultado (en la carpeta de salida):
  Productos-Eleventa.xlsx  -> Inventario > Importar Excel (ya existe en Compoint Punto)
  Clientes-Eleventa.xlsx   -> Clientes > Importar Excel

Uso (la base .FDB es de Firebird 2.5, que es la que usa Eleventa):
  python3 convertir_fdb.py PDVDATA.FDB -o salida --lib /ruta/libfbembed.so.2.5.4
  python3 convertir_fdb.py PDVDATA.FDB -o salida --servidor localhost --clave masterkey
Requiere:  pip install fdb openpyxl
"""
import argparse
import os
import sys

import fdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SIN_DEPARTAMENTO = '- Sin Departamento -'


def octetos(columna, largo):
    # Los textos de Eleventa vienen en Windows-1252 pero la base no lo declara;
    # se piden como bytes y se leen a mano para no perder acentos ni eñes.
    return f'cast({columna} as varchar({largo}) character set octets)'


def texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, bytes):
        return valor.decode('cp1252', errors='replace').strip()
    return str(valor).strip()


def numero(valor, decimales=2):
    if valor is None:
        return 0.0
    return round(float(valor), decimales)


def conectar(args):
    if args.lib:
        fdb.load_api(args.lib)
    if args.servidor:
        return fdb.connect(host=args.servidor, database=os.path.abspath(args.fdb),
                           user=args.usuario, password=args.clave, charset='NONE')
    # Modo "embebido": abre el archivo directo, sin usuario ni clave.
    return fdb.connect(dsn=os.path.abspath(args.fdb), charset='NONE')


def leer_productos(cur):
    cur.execute(f"""
        select {octetos('p.codigo', 50)}, {octetos('p.descripcion', 255)}, p.tventa,
               p.pcosto, p.pventa, p.pfinal, p.mayoreo, p.pmayoreofinal,
               p.dinventario, p.dinvminimo, p.dinvmaximo, p.usa_inventario, p.es_kit,
               {octetos('d.nombre', 255)}
        from productos p
        left join departamentos d on d.id = p.dept
        where p.eliminado_en is null
        order by p.descripcion
    """)
    filas, avisos = [], {'sin_nombre': 0, 'con_impuestos': 0, 'kits': 0, 'sin_inventario': 0, 'negativos': 0}
    for (codigo, nombre, tventa, pcosto, pventa, pfinal, mayoreo, pmayfinal,
         inv, inv_min, inv_max, usa_inv, es_kit, depto) in cur.fetchall():
        nombre = texto(nombre)
        if not nombre:
            avisos['sin_nombre'] += 1
            continue
        # Precio que paga el cliente en caja: el "final" (con impuestos) si existe.
        precio = numero(pfinal) or numero(pventa)
        if numero(pfinal) and abs(numero(pfinal) - numero(pventa)) > 0.004:
            avisos['con_impuestos'] += 1
        mayoreo_final = numero(pmayfinal) or numero(mayoreo)
        usa_inv = str(usa_inv).lower() in ('t', '1', 'true')
        if str(es_kit).lower() in ('t', '1', 'true'):
            avisos['kits'] += 1
        if not usa_inv:
            avisos['sin_inventario'] += 1
        existencia = numero(inv, 3) if usa_inv else 0
        if existencia < 0:
            avisos['negativos'] += 1
        depto = texto(depto)
        filas.append([
            texto(codigo),
            nombre,
            numero(pcosto),
            precio,
            mayoreo_final or None,
            existencia,
            numero(inv_min, 3),
            numero(inv_max, 3) or None,
            '' if depto == SIN_DEPARTAMENTO else depto,
            # En Eleventa 'D' (decimal) es producto a granel; 'U' es por unidad.
            'A Granel' if str(tventa).strip().upper() == 'D' else 'Por Pieza',
        ])
    return filas, avisos


def leer_clientes(cur):
    cur.execute(f"""
        select {octetos('c.nombres', 255)}, {octetos('c.apellidos', 255)}, {octetos('c.telefono', 50)},
               cr.tiene_credito, cr.limite_credito, cr.saldo_actual
        from clientesv2 c
        left join clientesv2_credito cr on cr.clientesv2_id = c.id and cr.eliminado_en is null
        where c.de_sistema = 0 and c.activo = 1
        order by c.nombres, c.apellidos
    """)
    filas = []
    for nombres, apellidos, telefono, tiene_credito, limite, saldo in cur.fetchall():
        nombre = (texto(nombres) + ' ' + texto(apellidos)).strip()
        if not nombre:
            continue
        con_credito = bool(tiene_credito)
        filas.append([
            nombre,
            texto(telefono),
            numero(limite) if con_credito else 0,
            max(numero(saldo), 0) if con_credito else 0,
        ])
    return filas


def escribir_excel(ruta, hoja, encabezados, filas, anchos, formatos_dinero=(), texto_cols=()):
    libro = Workbook()
    ws = libro.active
    ws.title = hoja
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='2D6B34')
        celda.alignment = Alignment(horizontal='center', vertical='center')
    for fila in filas:
        ws.append(fila)
    for r in range(2, ws.max_row + 1):
        for c in texto_cols:                      # códigos y teléfonos como texto (no pierden ceros)
            celda = ws.cell(row=r, column=c)
            celda.value = '' if celda.value is None else str(celda.value)
            celda.number_format = '@'
        for c in formatos_dinero:
            ws.cell(row=r, column=c).number_format = '#,##0.00'
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    ws.freeze_panes = 'A2'
    libro.save(ruta)


def main():
    ap = argparse.ArgumentParser(description='Eleventa (.FDB) -> Excel para Compoint Punto')
    ap.add_argument('fdb', help='archivo PDVDATA.FDB de Eleventa')
    ap.add_argument('-o', '--salida', default='.', help='carpeta donde se guardan los Excel')
    ap.add_argument('--lib', help='ruta de libfbembed.so (Firebird 2.5 embebido)')
    ap.add_argument('--servidor', help='si Firebird corre como servidor, su dirección (ej. localhost)')
    ap.add_argument('--usuario', default='SYSDBA')
    ap.add_argument('--clave', default='masterkey')
    args = ap.parse_args()

    os.makedirs(args.salida, exist_ok=True)
    con = conectar(args)
    cur = con.cursor()

    productos, avisos = leer_productos(cur)
    clientes = leer_clientes(cur)
    con.close()

    print(f'Productos encontrados: {len(productos)}')
    if productos:
        ruta = os.path.join(args.salida, 'Productos-Eleventa.xlsx')
        escribir_excel(
            ruta, 'Productos',
            ['Código', 'Producto', 'P. Costo', 'P. Venta', 'P. Mayoreo', 'Existencia',
             'Inv. Mínimo', 'Inv. Máximo', 'Departamento', 'Tipo de venta'],
            productos, [18, 44, 12, 12, 12, 12, 12, 12, 22, 14],
            formatos_dinero=(3, 4, 5), texto_cols=(1,))
        print('  ->', ruta)
    print(f'Clientes encontrados: {len(clientes)}  (no cuentan "Público en General" ni "Venta al extranjero")')
    if clientes:
        ruta = os.path.join(args.salida, 'Clientes-Eleventa.xlsx')
        escribir_excel(
            ruta, 'Clientes', ['Nombre', 'Teléfono', 'Límite de crédito', 'Saldo'],
            clientes, [38, 18, 18, 14], formatos_dinero=(3, 4), texto_cols=(2,))
        print('  ->', ruta)

    print('Avisos:')
    print(f"  Sin nombre (se omitieron): {avisos['sin_nombre']}")
    print(f"  Con impuestos (se usó el precio final): {avisos['con_impuestos']}")
    print(f"  Que no controlaban inventario (quedan en 0): {avisos['sin_inventario']}")
    print(f"  Kits (se pasan como producto normal): {avisos['kits']}")
    print(f"  Con existencia negativa: {avisos['negativos']}")


if __name__ == '__main__':
    sys.exit(main())
